"""Read-only adapters. Source text is data, never executable instructions."""
import csv
import io
import json
from zipfile import BadZipFile
from datetime import date, datetime
from pydantic import BaseModel, ConfigDict, ValidationError
from typing import Literal
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from ...auth import Problem

TEMPLATES = {
    'csv-v1': {'record_id': 'record_id', 'report_id': 'report_id', 'project_id': 'project_id',
               'description': 'description', 'record_kind': 'record_kind', 'occurred_on': 'inspection_date',
               'images': 'images', 'original_index': 'original_index'},
    'excel-v1': {'record_id': '来源记录ID', 'report_id': '报告ID', 'project_id': '项目ID',
                 'description': '隐患描述', 'record_kind': '记录类型', 'occurred_on': '检查日期',
                 'images': '图片关系', 'original_index': '原始序号'},
    'governance-v1': {'record_id': 'record_id', 'report_id': 'report_id', 'project_id': 'project_id',
                      'description': 'description', 'record_kind': 'record_kind', 'occurred_on': 'inspection_date',
                      'images': 'images', 'original_index': 'original_index'},
}
DEFAULT_TEMPLATE = {'csv': 'csv-v1', 'excel': 'excel-v1', 'governance_export': 'governance-v1'}
MIME = {'csv': 'text/csv', 'excel': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'governance_export': 'application/json'}


class FieldValidation(BaseModel):
    model_config = ConfigDict(strict=True)
    description: str
    record_kind: Literal['detail', 'review', 'summary']
    occurred_on: date | None


def read_rows(kind, content):
    if len(content) > 5_000_000:
        raise Problem(422, 'INPUT_TOO_LARGE', '首版输入上限5MB，请按批次接入。')
    try:
        if kind == 'csv':
            reader = csv.DictReader(io.StringIO(content.decode('utf-8-sig')))
            headers = reader.fieldnames or []
            if not headers or len(headers) != len(set(headers)):
                raise ValueError('duplicate or missing headers')
            return [(n, dict(row), {}, {'line': n}) for n, row in enumerate(reader, 2)]
        if kind == 'excel':
            book = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
            try:
                sheet = book.worksheets[0]
                values = iter(sheet.values)
                headers = [str(v) if v is not None else '' for v in next(values)]
                if not all(headers) or len(set(headers)) != len(headers):
                    raise ValueError('invalid headers')
                def cell(v):
                    return v.isoformat() if isinstance(v, (date, datetime)) else v
                return [(n, dict(zip(headers, map(cell, row))), {},
                         {'sheet': sheet.title, 'row': n, 'range': f'A{n}:{get_column_letter(sheet.max_column)}{n}'})
                        for n, row in enumerate(values, 2) if any(v is not None for v in row)]
            finally:
                book.close()
        exported = json.loads(content)
        if exported.get('format') != 'governance-export-v1' or not isinstance(exported.get('records'), list):
            raise ValueError('unsupported export')
        return [(n, dict(row['raw_fields']), {'effective_fields': row.get('effective_fields', {}),
                'revision': row.get('revision'), 'modified_by': row.get('modified_by'),
                'original_source_id': row.get('source_id'), 'field_provenance': row.get('field_provenance', {})},
                 {'export_index': n}) for n, row in enumerate(exported['records'], 1)]
    except (ValueError, TypeError, KeyError, UnicodeError, StopIteration, OSError, BadZipFile, InvalidFileException) as exc:
        raise Problem(422, 'INPUT_FORMAT_ERROR', '输入格式或表头不符合接入器要求。') from exc


def normalize(raw, governed, mapping, project):
    fields = {key: raw.get(column) for key, column in mapping.items()}
    issues = []
    errors = []
    if fields.get('project_id') != project:
        errors.append('项目范围不匹配，未接入')
    external = str(fields.get('record_id') or '').strip()
    report = str(fields.get('report_id') or '').strip()
    if not report:
        errors.append('缺少稳定报告ID')
    if not external:
        original = fields.get('original_index')
        if report and original is not None and str(original).strip():
            external = f'fallback:{project}:{report}:{original}'
            issues.append('使用报告ID+项目+来源原始序号作为后备身份；不能使用当前行号或可变描述')
        else:
            errors.append('缺少稳定记录ID且无来源原始序号，须人工补充身份')
    if len(external) > 256 or len(report) > 256:
        errors.append('来源标识过长')
    effective = dict(fields)
    # Only explicit export fields override raw values; source identity/scope stay server-controlled.
    for key, value in governed.get('effective_fields', {}).items():
        if key not in {'record_id', 'report_id', 'project_id', 'original_index'}:
            effective[key] = value
    if effective.get('record_kind') not in ('detail', 'review', 'summary'):
        errors.append('记录类型须为detail/review/summary')
    raw_date = effective.get('occurred_on')
    when = None
    if raw_date:
        try:
            when = date.fromisoformat(str(raw_date)[:10])
        except ValueError:
            issues.append('日期无效，原值保留，生效日期置空')
    description = effective.get('description')
    if not isinstance(description, str) or not description.strip():
        issues.append('缺少有效描述，保留记录供核查')
        description = ''
    try:
        validated = FieldValidation(description=description, record_kind=effective.get('record_kind'), occurred_on=when)
        effective.update(validated.model_dump(mode='json'))
    except ValidationError:
        errors.append('字段类型校验失败')
    images = effective.get('images') or []
    if isinstance(images, str):
        try:
            images = json.loads(images)
        except ValueError:
            images = []
            issues.append('images不是有效JSON列表，坏值已保留在原始字段')
    clean_images = []
    if not isinstance(images, list):
        issues.append('images类型错误，应为列表')
    else:
        for image in images:
            if not isinstance(image, dict) or image.get('type') not in {'site_photo', 'text_screenshot'}:
                issues.append('图片关系字段无效，原值保留')
                continue
            confidence = image.get('relation_confidence', 'candidate')
            if confidence not in {'anchored', 'candidate'}:
                confidence = 'candidate'
            clean_images.append({**image, 'relation_confidence': confidence})
    effective['images'] = clean_images
    return {'external_id': external, 'report_id': report, 'effective': effective,
            'issues': issues, 'errors': errors, 'governed': governed}
